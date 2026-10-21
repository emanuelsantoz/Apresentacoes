"""
Monta exports/mvp-conferencia-contabil.pptx a partir de exports/slides/*.png
e exports/notes.json (gerados por `npm run export:pdf`).
Uso: pip install python-pptx && python scripts/export-pptx.py
"""
import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
EXPORTS = ROOT / "exports"
OUT = EXPORTS / "mvp-conferencia-contabil.pptx"

notes = json.loads((EXPORTS / "notes.json").read_text(encoding="utf-8"))
images = sorted((EXPORTS / "slides").glob("slide-*.png"))

prs = Presentation()
prs.slide_width = Emu(12192000)  # 13,333 in — 16:9
prs.slide_height = Emu(6858000)  # 7,5 in
blank = prs.slide_layouts[6]

for img, meta in zip(images, notes):
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
    slide.notes_slide.notes_text_frame.text = f"{meta['n']:02d} · {meta['title']}\n\n{meta['notes']}"

prs.save(OUT)
print(f"✓ {OUT.relative_to(ROOT)} ({len(images)} slides)")
